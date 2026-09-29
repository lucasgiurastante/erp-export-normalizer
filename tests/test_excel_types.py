"""Excel output: native cell types instead of text."""

from __future__ import annotations

import contextlib
import datetime
import decimal
import io
import os
import tempfile
import unittest

from cli import main
from core.schema import build_schema
from core.validator import FieldValue, RecordResult
from core.writer import ExcelWriter

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCHEMA = os.path.join(REPO_ROOT, "core", "formats", "jde_ar.yaml")
FIXTURE = os.path.join(REPO_ROOT, "examples", "data", "jde_ar.txt")
PACKED = os.path.join(REPO_ROOT, "core", "formats", "cobol_packed.yaml")
PACKED_FIXTURE = os.path.join(REPO_ROOT, "examples", "data", "cobol_packed.txt")

SCHEMA_OBJ = build_schema(
    {
        "format": "t",
        "version": "1.0.0",
        "record_length": 24,
        "fields": [
            {"name": "s", "start": 0, "length": 6},
            {
                "name": "d",
                "start": 6,
                "length": 8,
                "type": "date",
                "format": "YYYYMMDD",
            },
            {"name": "n", "start": 14, "length": 10, "type": "decimal", "scale": 2},
        ],
    }
)


def result(values: list[tuple[str, object]]) -> RecordResult:
    return RecordResult(
        line=1,
        ok=True,
        errors=[],
        fields=[FieldValue(n, v, "") for n, v in values],
    )


def read_back(path: str):
    from openpyxl import load_workbook

    return load_workbook(path).active


class TestExcelNativeTypes(unittest.TestCase):
    def _cell(self, res: RecordResult, index: int, **kwargs):
        """Write one record and return the cell at `index` of its row.

        The workbook is read inside the temp directory: handing back a path
        that is about to be deleted would make every test fail the same way.
        """
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "o.xlsx")
            w = ExcelWriter(SCHEMA_OBJ, path, **kwargs)
            w.write(res)
            w.finish()
            ws = load_workbook(path).active
            return list(ws.iter_rows(min_row=2, max_row=2))[0][index]

    def test_date_is_a_date_cell(self):
        cell = self._cell(
            result([("s", "a"), ("d", "2024-08-15"), ("n", decimal.Decimal("1"))]), 1
        )
        self.assertIsInstance(cell.value, (datetime.date, datetime.datetime))
        self.assertEqual(cell.data_type, "d")

    def test_decimal_is_a_numeric_cell(self):
        cell = self._cell(
            result(
                [("s", "a"), ("d", "2024-08-15"), ("n", decimal.Decimal("1234.56"))]
            ),
            2,
        )
        self.assertIsInstance(cell.value, (int, float))
        self.assertNotIsInstance(cell.value, str)
        self.assertEqual(cell.data_type, "n")
        self.assertEqual(cell.value, 1234.56)

    def test_amount_is_summable_in_excel(self):
        """The point of native types: SUM() has to work on the column."""
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "o.xlsx")
            w = ExcelWriter(SCHEMA_OBJ, path)
            for amount in ("10.10", "20.20", "30.30"):
                w.write(
                    result(
                        [
                            ("s", "a"),
                            ("d", "2024-08-15"),
                            ("n", decimal.Decimal(amount)),
                        ]
                    )
                )
            w.finish()
            ws = read_back(path)
            total = sum(row[2].value for row in ws.iter_rows(min_row=2))
            self.assertAlmostEqual(total, 60.60, places=2)

    def test_none_is_a_blank_cell_not_text(self):
        cell = self._cell(result([("s", "a"), ("d", "2024-08-15"), ("n", None)]), 2)
        self.assertIsNone(cell.value)

    def test_non_iso_date_stays_text(self):
        """A masked value must not turn redaction into a hard failure."""
        cell = self._cell(
            result([("s", "a"), ("d", "20********"), ("n", decimal.Decimal("1"))]), 1
        )
        self.assertEqual(cell.value, "20********")
        self.assertIsInstance(cell.value, str)

    def test_float_decimals_is_configurable(self):
        cell = self._cell(
            result(
                [("s", "a"), ("d", "2024-08-15"), ("n", decimal.Decimal("1.23456789"))]
            ),
            2,
            float_decimals=4,
        )
        self.assertEqual(cell.value, 1.2346)

    def test_negative_float_decimals_rejected(self):
        from openpyxl import Workbook  # noqa: F401  (ensure the dep is present)

        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                ExcelWriter(SCHEMA_OBJ, os.path.join(tmp, "o.xlsx"), float_decimals=-1)

    def test_header_row_is_names(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "o.xlsx")
            w = ExcelWriter(SCHEMA_OBJ, path)
            w.write(
                result([("s", "a"), ("d", "2024-08-15"), ("n", decimal.Decimal("1"))])
            )
            w.finish()
            ws = read_back(path)
            header = [c.value for c in list(ws.iter_rows(min_row=1, max_row=1))[0]]
            self.assertEqual(header, ["s", "d", "n"])


class TestExcelCli(unittest.TestCase):
    def _rows(self, schema: str, src: str, *extra: str):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "o.xlsx")
            with (
                contextlib.redirect_stderr(io.StringIO()),
                contextlib.redirect_stdout(io.StringIO()),
            ):
                code = main(
                    [
                        "--schema",
                        schema,
                        "--input",
                        src,
                        "--output",
                        out,
                        "--format",
                        "excel",
                        *extra,
                    ]
                )
            # the packed fixture carries one intentional bad record
            self.assertIn(code, (0, 3))
            ws = read_back(out)
            return list(ws.iter_rows(values_only=True))

    def test_jde_ar_types(self):
        rows = self._rows(SCHEMA, FIXTURE)
        self.assertEqual(rows[0], ("id", "type", "date", "amount", "currency"))
        self.assertEqual(rows[1][0], "AR1001")
        self.assertIsInstance(rows[1][2], (datetime.date, datetime.datetime))
        self.assertEqual(rows[1][3], 1234.56)
        self.assertEqual(rows[1][4], "USD")

    def test_packed_becomes_numeric(self):
        rows = self._rows(PACKED, PACKED_FIXTURE)
        self.assertEqual(rows[1][2], 1234567.99)
        self.assertIsInstance(rows[1][2], float)

    def test_excel_float_decimals_flag(self):
        rows = self._rows(SCHEMA, FIXTURE, "--excel-float-decimals", "0")
        self.assertEqual(rows[1][3], 1235)


if __name__ == "__main__":
    unittest.main()
